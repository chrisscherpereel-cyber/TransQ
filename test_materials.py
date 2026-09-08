"""Reading the lecture's own slides, and steering questions by what will be tested.

Two additions with one motive: a transcript is a poor record of what was *shown*.
The term printed on a slide arrives in the audio as "this thing here", a table of
figures is never read aloud, and a gesture at a diagram transcribes as silence.
The deck restores that. The exam-topic list restores the other missing half —
what the lecture was *for* — which the app could otherwise only infer from a
summary that is itself a model's guess.

Real files are built here rather than stubbed, because the failure that matters
is "python-pptx returned something unexpected", and a stub cannot fail that way.
"""

from __future__ import annotations

import io

import pytest

from src.chunking import EXAM_TOPIC_WEIGHT, allocate_by_importance, chunk_importance
from src.materials import (
    MIN_SECTION_CHARS,
    Material,
    MaterialError,
    extract_material,
    material_context,
    relevant_sections,
)
from src.mcq import _build_exam_clause
from src.schema import Chunk, Summary

CHASE = (
    "Chase strategy tracks demand period by period through hiring and firing, "
    "trading workforce stability for low inventory."
)
LEVEL = (
    "Level strategy holds the workforce steady and absorbs demand variation "
    "with inventory buffers and backlogs."
)


# --------------------------------------------------------------------------- #
# Building real files
# --------------------------------------------------------------------------- #


def make_pptx(slides: list[tuple[str, str, str]]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    for title, body, notes in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[5])  # title only
        slide.shapes.title.text = title
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(3))
        box.text_frame.text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def make_pdf(pages: list[str]) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)
    for page in pages:
        y = 720
        for line in page.splitlines():
            pdf.drawString(72, y, line)
            y -= 16
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_docx(blocks: list[tuple[str, str]]) -> bytes:
    import docx

    document = docx.Document()
    for heading, body in blocks:
        if heading:
            document.add_heading(heading, level=1)
        document.add_paragraph(body)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# Extraction
# --------------------------------------------------------------------------- #


def test_a_powerpoint_deck_becomes_titled_slides():
    blob = make_pptx(
        [
            ("Chase strategy", CHASE, ""),
            ("Level strategy", LEVEL, "The point students miss is the buffer cost."),
        ]
    )
    material = extract_material("week4.pptx", blob)

    assert material.kind == "PowerPoint"
    assert material.section_noun == "slide"
    assert [s.title for s in material.sections] == ["Chase strategy", "Level strategy"]
    assert "hiring and firing" in material.sections[0].text


def test_speaker_notes_are_read():
    """Often where the argument lives, and invisible in the room — so exactly
    the material the transcript is least likely to duplicate."""
    blob = make_pptx([("Level strategy", LEVEL, "Students miss the buffer cost.")])
    material = extract_material("w.pptx", blob)

    assert "buffer cost" in material.sections[0].full_text
    assert "Speaker notes:" in material.sections[0].full_text


def test_a_pdf_becomes_pages():
    material = extract_material("handout.pdf", make_pdf([CHASE, LEVEL]))
    assert material.kind == "PDF"
    assert material.section_noun == "page"
    assert len(material.sections) == 2
    assert "Chase strategy" in material.text


def test_a_word_document_splits_on_headings():
    material = extract_material(
        "notes.docx", make_docx([("Chase strategy", CHASE), ("Level strategy", LEVEL)])
    )
    assert [s.title for s in material.sections] == ["Chase strategy", "Level strategy"]


def test_markdown_splits_on_headings():
    text = f"# Chase strategy\n{CHASE}\n\n# Level strategy\n{LEVEL}\n".encode()
    material = extract_material("outline.md", text)
    assert [s.title for s in material.sections] == ["Chase strategy", "Level strategy"]


def test_empty_slides_are_dropped_not_carried():
    """A "Questions?" slide matches nothing and dilutes every relevance score."""
    blob = make_pptx([("Chase strategy", CHASE, ""), ("Questions?", "", "")])
    material = extract_material("w.pptx", blob)

    assert len(material.sections) == 1
    assert all(len(s.full_text) >= MIN_SECTION_CHARS for s in material.sections)


def test_sections_are_renumbered_after_the_empties_go():
    blob = make_pptx([("?", "", ""), ("Chase strategy", CHASE, "")])
    material = extract_material("w.pptx", blob)
    assert [s.index for s in material.sections] == [0]
    assert material.sections[0].label.startswith("1.")


# --------------------------------------------------------------------------- #
# A bad file costs the upload, never the lecture
# --------------------------------------------------------------------------- #


def test_an_unsupported_type_is_refused_by_name():
    with pytest.raises(MaterialError) as exc:
        extract_material("lecture.key", b"whatever")
    assert "not supported" in str(exc.value)


def test_a_corrupt_file_raises_the_one_type_callers_catch():
    """It must be MaterialError, not a pptx internal — the caller catches one
    type so an unreadable deck cannot take the transcript down with it."""
    with pytest.raises(MaterialError):
        extract_material("broken.pptx", b"this is not a zip archive")


def test_a_deck_of_images_says_what_is_wrong():
    with pytest.raises(MaterialError) as exc:
        extract_material("scans.md", b"   \n\n  \n")
    assert "OCR" in str(exc.value)


# --------------------------------------------------------------------------- #
# Matching material to a moment in the lecture
# --------------------------------------------------------------------------- #


@pytest.fixture
def deck() -> Material:
    return extract_material(
        "week4.pptx",
        make_pptx(
            [
                ("Chase strategy", CHASE, ""),
                ("Level strategy", LEVEL, ""),
                (
                    "Course admin",
                    "Office hours are Tuesday. The midterm covers weeks one to six.",
                    "",
                ),
            ]
        ),
    )


def test_a_window_gets_the_slides_that_belong_to_it(deck):
    spoken = (
        "so with chase we track demand each period, hiring and firing as we go, "
        "which trades workforce stability against carrying inventory"
    )
    matched = relevant_sections(spoken, deck, limit=2)
    assert matched
    assert matched[0].title == "Chase strategy"
    assert "Course admin" not in [s.title for s in matched]


def test_an_unrelated_window_gets_nothing_rather_than_the_least_bad_slide(deck):
    """Handing every window its three closest slides invites questions about
    material the instructor had not reached yet."""
    assert relevant_sections("good morning everyone, can you all hear me", deck) == []


def test_matching_is_bounded(deck):
    assert len(relevant_sections(CHASE + " " + LEVEL, deck, limit=1)) == 1


def test_no_material_matches_nothing_without_raising():
    assert relevant_sections("anything", None) == []
    assert relevant_sections("anything", Material()) == []
    assert relevant_sections("", None) == []


def test_the_material_prompt_says_why_the_slides_are_there(deck):
    clause = material_context(deck.sections[:1], deck.section_noun)
    assert "Chase strategy" in clause
    assert "this" in clause, "it should explain the pronoun problem it solves"
    assert material_context([], "slide") == ""


# --------------------------------------------------------------------------- #
# Round-tripping through storage
# --------------------------------------------------------------------------- #


def test_a_deck_survives_being_saved_and_reloaded(deck):
    restored = Material.from_dict(deck.to_dict())
    assert restored.filename == deck.filename
    assert restored.kind == deck.kind
    assert [s.title for s in restored.sections] == [s.title for s in deck.sections]
    assert restored.text == deck.text


def test_a_malformed_stored_deck_degrades_to_empty():
    assert Material.from_dict({}).sections == []
    assert Material.from_dict({"sections": ["not a dict"]}).sections == []


# --------------------------------------------------------------------------- #
# Exam topics outrank everything the app infers
# --------------------------------------------------------------------------- #


def windows() -> list[Chunk]:
    return [
        Chunk(index=0, start=0, end=600, text="welcome syllabus office hours midterm dates"),
        Chunk(index=1, start=600, end=1200, text=CHASE * 3),
        Chunk(index=2, start=1200, end=1800, text=LEVEL * 3),
        Chunk(index=3, start=1800, end=2400, text="a story from my consulting days"),
    ]


def test_exam_topics_pull_questions_towards_themselves():
    counts = allocate_by_importance(
        10, windows(), None, ["level strategy inventory buffers backlogs"]
    )
    assert counts[2] == max(counts), "the examined window should carry the most"
    assert sum(counts) == 10


def test_exam_topics_outrank_the_summarys_own_judgment():
    """The summary is a model's account of what a lecture contained; this is the
    instructor's statement of what it was for. They can disagree."""
    summary = Summary(learning_objectives=["Explain chase strategy hiring and firing"])
    with_topics = chunk_importance(
        windows(), summary, ["level strategy inventory buffers"]
    )
    assert with_topics[2] > with_topics[1]
    assert EXAM_TOPIC_WEIGHT > 2.0


def test_no_exam_topics_leaves_the_old_behaviour_intact():
    summary = Summary(learning_objectives=["Explain chase strategy hiring and firing"])
    assert chunk_importance(windows(), summary, []) == chunk_importance(
        windows(), summary
    )


def test_blank_lines_in_the_topic_box_are_ignored():
    counts = allocate_by_importance(6, windows(), None, ["", "   ", "level strategy"])
    assert sum(counts) == 6


def test_the_exam_clause_is_emphatic_and_bounded():
    clause = _build_exam_clause(["Chase versus level", "Cost under volatility"])
    assert "EXAM" in clause
    assert "Chase versus level" in clause
    assert "inventing" in clause, "it must not invite fabrication when unsupported"

    assert _build_exam_clause([]) == ""
    assert _build_exam_clause(["", "  "]) == ""
    assert _build_exam_clause([f"Topic {i}" for i in range(40)]).count("\n- ") <= 20


# --------------------------------------------------------------------------- #
# Saved with the lecture
# --------------------------------------------------------------------------- #


def test_the_deck_and_topics_reopen_with_the_lecture(deck, tmp_path):
    """Re-uploading slides and retyping a topic list to regenerate one question
    set is exactly the friction that stops people regenerating."""
    from src.library import TranscriptLibrary
    from src.schema import Segment, Transcript
    from src.storage import build_store

    secret = "a-long-random-app-secret-for-tests"
    store, _ = build_store({"APP_SECRET": secret, "DATA_DIR": str(tmp_path)})
    transcript = Transcript(
        segments=[Segment(index=0, start=0, end=60, text=CHASE)], duration=600.0
    )

    entry = TranscriptLibrary(store, "chris").save(
        transcript,
        title="MGT 301 — Week 4",
        material=deck,
        exam_topics=["Chase versus level", "Cost under volatility"],
    )

    reopened, _ = build_store({"APP_SECRET": secret, "DATA_DIR": str(tmp_path)})
    saved = TranscriptLibrary(reopened, "chris").load(entry.id)

    assert saved.material is not None
    assert saved.material.filename == "week4.pptx"
    assert [s.title for s in saved.material.sections] == [
        s.title for s in deck.sections
    ]
    assert saved.exam_topics == ["Chase versus level", "Cost under volatility"]


def test_the_listing_shows_that_a_lecture_has_slides(deck, tmp_path):
    from src.library import TranscriptLibrary
    from src.schema import Segment, Transcript
    from src.storage import GuardedStore, MemoryStore

    library = TranscriptLibrary(GuardedStore(MemoryStore()), "chris")
    library.save(
        Transcript(segments=[Segment(index=0, start=0, end=60, text=CHASE)]),
        material=deck,
        exam_topics=["one", "two"],
    )
    listed = library.entries()[0]
    assert listed.material_filename == "week4.pptx"
    assert listed.exam_topics == 2


def test_a_lecture_saved_without_either_still_loads(tmp_path):
    """Everything saved before this feature existed must reopen unchanged."""
    from src.library import TranscriptLibrary
    from src.schema import Segment, Transcript
    from src.storage import GuardedStore, MemoryStore

    library = TranscriptLibrary(GuardedStore(MemoryStore()), "chris")
    entry = library.save(
        Transcript(segments=[Segment(index=0, start=0, end=60, text=CHASE)])
    )
    saved = library.load(entry.id)

    assert saved.material is None
    assert saved.exam_topics == []
